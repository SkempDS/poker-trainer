# -*- coding: utf-8 -*-
"""python-docx 中文排版公用件。"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_CN = "微软雅黑"
FONT_EN = "Segoe UI"


def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT_EN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.45
    for lvl, size, color in ((1, 16, "1F3864"), (2, 13, "2E5496"), (3, 11.5, "333333")):
        s = doc.styles[f"Heading {lvl}"]
        s.font.name = FONT_EN
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        s.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        s.paragraph_format.space_after = Pt(6)
    return doc


def cn(run, font=FONT_CN):
    """给 run 打中文字体，缺 rFonts 时补上。"""
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:eastAsia"), font)
    return run


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def cell_text(cell, text, size=9, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT_EN
    r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def set_borders(table, color="BFBFBF", sz=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def repeat_header(table):
    """跨页时重复首行表头，并禁止单行被拆开。"""
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        if i == 0:
            th = OxmlElement("w:tblHeader")
            trPr.append(th)


def make_table(doc, data, widths=None, header=True, size=9, header_fill="1F3864",
               header_color="FFFFFF", zebra="F2F5FA", align=WD_ALIGN_PARAGRAPH.CENTER,
               col_align=None):
    rows = len(data)
    cols = len(data[0])
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_borders(t)
    repeat_header(t)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            a = align
            if col_align and j < len(col_align) and col_align[j] is not None:
                a = col_align[j]
            if i == 0 and header:
                cell_text(c, val, size=size, bold=True, color=header_color, align=a)
                shade(c, header_fill)
            else:
                cell_text(c, val, size=size, align=a)
                if zebra and i % 2 == 0:
                    shade(c, zebra)
            if widths:
                c.width = Cm(widths[j])
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    return t


def para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=6, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = FONT_EN
    r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def rich(doc, parts, size=10.5, space_after=6):
    """parts = [(text, bold, color), ...] 同段混排。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    for text, bold, color in parts:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT_EN
        r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.3)
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.name = FONT_EN
    r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    r.font.color.rgb = RGBColor.from_string("595959")
    return p
